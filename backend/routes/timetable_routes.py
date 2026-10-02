import io
from flask import Blueprint, request, jsonify, send_file
import models.timetable as TimetableModel
from config import config

timetable_bp = Blueprint("timetable", __name__)

DAYS  = config.DAYS
SLOTS = config.SLOTS


@timetable_bp.route("/", methods=["GET"])
def get_timetables():
    tts = TimetableModel.get_all()
    return jsonify(tts), 200


@timetable_bp.route("/<int:timetable_id>", methods=["GET"])
def get_timetable(timetable_id):
    tt = TimetableModel.get_by_id(timetable_id)
    if not tt:
        return jsonify({"error": "Timetable not found"}), 404
    entries = TimetableModel.get_entries(timetable_id)
    return jsonify({**tt, "entries": entries}), 200


@timetable_bp.route("/<int:timetable_id>/status", methods=["PUT"])
def update_status(timetable_id):
    data   = request.get_json()
    status = data.get("status", "saved")
    if status not in ("generated", "saved", "archived"):
        return jsonify({"error": "Invalid status"}), 400
    TimetableModel.update_status(timetable_id, status)
    return jsonify({"message": "Status updated"}), 200


@timetable_bp.route("/<int:timetable_id>", methods=["DELETE"])
def delete_timetable(timetable_id):
    if not TimetableModel.get_by_id(timetable_id):
        return jsonify({"error": "Timetable not found"}), 404
    TimetableModel.delete(timetable_id)
    return jsonify({"message": "Timetable deleted"}), 200


@timetable_bp.route("/stats", methods=["GET"])
def get_stats():
    return jsonify(TimetableModel.get_stats()), 200


# ── Export PDF ────────────────────────────────────────────────────────────────

@timetable_bp.route("/<int:timetable_id>/export/pdf", methods=["GET"])
def export_pdf(timetable_id):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    tt = TimetableModel.get_by_id(timetable_id)
    if not tt:
        return jsonify({"error": "Timetable not found"}), 404

    entries = TimetableModel.get_entries(timetable_id)
    buf     = io.BytesIO()
    doc     = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 leftMargin=1*cm, rightMargin=1*cm,
                                 topMargin=1.5*cm, bottomMargin=1*cm)
    styles  = getSampleStyleSheet()
    story   = []

    # Group entries by division
    div_map = {}
    for e in entries:
        div_map.setdefault(e["division_name"], []).append(e)

    for div_name, div_entries in sorted(div_map.items()):
        story.append(Paragraph(f"{tt['class_name']} – Division {div_name}", styles["Heading2"]))
        story.append(Spacer(1, 0.3*cm))

        # Build grid
        header = ["TIME"] + DAYS
        rows   = [header]

        for slot_info in SLOTS:
            row = [slot_info["label"]]
            for d_idx, _ in enumerate(DAYS):
                cell = ""
                for e in div_entries:
                    if e["day"] == d_idx and e["slot"] == slot_info["index"]:
                        cell = f"{e['subject_name']}\n{e['teacher_name']}\n{e['classroom_name']}"
                row.append(cell)
            rows.append(row)

            # Insert break rows
            if slot_info["index"] == 0:
                rows.append(["12:00–12:40\nBREAK"] + [""] * 5)
            elif slot_info["index"] == 1:
                rows.append(["2:40–2:50\nBREAK"] + [""] * 5)

        t = Table(rows, colWidths=[3*cm] + [4.5*cm]*5)
        t.setStyle(TableStyle([
            ("BACKGROUND",   (0, 0), (-1, 0), colors.HexColor("#1e40af")),
            ("TEXTCOLOR",    (0, 0), (-1, 0), colors.white),
            ("FONTNAME",     (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",     (0, 0), (-1, -1), 8),
            ("ALIGN",        (0, 0), (-1, -1), "CENTER"),
            ("VALIGN",       (0, 0), (-1, -1), "MIDDLE"),
            ("GRID",         (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#f1f5f9")]),
        ]))
        story.append(t)
        story.append(Spacer(1, 0.8*cm))

    doc.build(story)
    buf.seek(0)
    return send_file(buf, mimetype="application/pdf",
                     as_attachment=True,
                     download_name=f"timetable_{timetable_id}.pdf")


# ── Export Excel ──────────────────────────────────────────────────────────────

@timetable_bp.route("/<int:timetable_id>/export/excel", methods=["GET"])
def export_excel(timetable_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    tt = TimetableModel.get_by_id(timetable_id)
    if not tt:
        return jsonify({"error": "Timetable not found"}), 404

    entries = TimetableModel.get_entries(timetable_id)
    wb      = Workbook()
    wb.remove(wb.active)  # remove default sheet

    div_map = {}
    for e in entries:
        div_map.setdefault(e["division_name"], []).append(e)

    header_fill = PatternFill("solid", fgColor="1E40AF")
    break_fill  = PatternFill("solid", fgColor="E2E8F0")
    bold_white  = Font(bold=True, color="FFFFFF")
    bold_dark   = Font(bold=True)
    center      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin        = Side(style="thin")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    for div_name, div_entries in sorted(div_map.items()):
        ws = wb.create_sheet(title=f"Division {div_name}")
        ws.column_dimensions["A"].width = 16
        for col in "BCDEF":
            ws.column_dimensions[col].width = 24

        # Header row
        header = ["TIME"] + DAYS
        for col_idx, val in enumerate(header, 1):
            cell = ws.cell(row=1, column=col_idx, value=val)
            cell.fill   = header_fill
            cell.font   = bold_white
            cell.alignment = center
            cell.border = border

        row_num = 2
        for slot_info in SLOTS:
            ws.cell(row=row_num, column=1, value=slot_info["label"]).alignment = center
            ws.cell(row=row_num, column=1).font   = bold_dark
            ws.cell(row=row_num, column=1).border = border
            ws.row_dimensions[row_num].height = 60

            for d_idx, _ in enumerate(DAYS):
                cell_val = ""
                for e in div_entries:
                    if e["day"] == d_idx and e["slot"] == slot_info["index"]:
                        cell_val = f"{e['subject_name']}\n{e['teacher_name']}\n{e['classroom_name']}"
                c = ws.cell(row=row_num, column=d_idx + 2, value=cell_val)
                c.alignment = center
                c.border    = border

            row_num += 1

            # Insert break row
            if slot_info["index"] in (0, 1):
                break_label = "12:00–12:40  BREAK" if slot_info["index"] == 0 else "2:40–2:50  BREAK"
                for col_idx in range(1, 7):
                    c = ws.cell(row=row_num, column=col_idx, value=break_label if col_idx == 1 else "")
                    c.fill      = break_fill
                    c.alignment = center
                    c.border    = border
                    c.font      = Font(italic=True, color="64748B")
                row_num += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=f"timetable_{timetable_id}.xlsx")
